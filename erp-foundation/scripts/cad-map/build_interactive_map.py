#!/usr/bin/env python3
"""Build a clean, interactive SVG map from the converted Bhachunda CAD files.

The CAD conversion provides two useful inputs:
  * a minJSON export, which preserves layers, label locations and polygon points;
  * an SVG export, which preserves the CAD drawing paths for the browser.

This script joins them by object index, associates every main survey label with
the polygon that contains it, and emits an SVG plus a reviewable JSON manifest.
It intentionally does not write to Supabase: the JSON is a staging artefact
that must be reviewed before parcel_map_features is imported.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SVG_NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", SVG_NS)

GUJARATI_DIGITS = str.maketrans("૦૧૨૩૪૫૬૭૮૯", "0123456789")
PRIMARY_POLYGON_LAYER = "0"
PRIMARY_LABEL_LAYER = "NEW SVY NO"
SECONDARY_POLYGON_LAYER = "Survey_Limit"
SECONDARY_LABEL_LAYER = "Survey No New"


@dataclass(frozen=True)
class Polygon:
    object_index: int
    points: tuple[tuple[float, float], ...]
    bbox: tuple[float, float, float, float]
    area: float


@dataclass(frozen=True)
class Label:
    object_index: int
    text: str
    normalized: str
    point: tuple[float, float]


def normalized_survey(value: Any) -> str:
    """Use the same Gujarati-digit/whitespace normalisation as Supabase."""

    return re.sub(r"\s+", "", str(value or "").strip().translate(GUJARATI_DIGITS).lower())


def survey_base(value: str) -> str:
    """Return the numeric root used only as a review candidate, never as proof."""

    match = re.match(r"\d+", value)
    return match.group(0) if match else value


def polygon_area(points: tuple[tuple[float, float], ...]) -> float:
    return abs(
        sum(
            x1 * y2 - x2 * y1
            for (x1, y1), (x2, y2) in zip(points, points[1:] + points[:1])
        )
        / 2
    )


def contains_point(point: tuple[float, float], points: tuple[tuple[float, float], ...]) -> bool:
    """Ray-casting point-in-polygon test. Survey labels are not on boundaries."""

    x, y = point
    inside = False
    for (x1, y1), (x2, y2) in zip(points, points[1:] + points[:1]):
        if (y1 > y) != (y2 > y):
            x_at_y = (x2 - x1) * (y - y1) / (y2 - y1) + x1
            if x < x_at_y:
                inside = not inside
    return inside


def layer_name_lookup(objects: list[dict[str, Any]]) -> dict[int, str]:
    return {
        item["handle"][-1]: item["name"]
        for item in objects
        if item.get("object") == "LAYER" and item.get("handle") and item.get("name")
    }


def object_layer(item: dict[str, Any], lookup: dict[int, str]) -> str | None:
    layer = item.get("layer")
    return lookup.get(layer[-1]) if layer else None


def collect_polygons(
    objects: list[dict[str, Any]], lookup: dict[int, str], layer: str
) -> list[Polygon]:
    polygons: list[Polygon] = []
    for item in objects:
        if (
            item.get("entity") != "LWPOLYLINE"
            or object_layer(item, lookup) != layer
            or item.get("entmode") != 2
        ):
            continue
        points = tuple(tuple(point[:2]) for point in item.get("points", []))
        if len(points) < 3:
            continue
        xs, ys = zip(*points)
        polygons.append(
            Polygon(
                object_index=item["index"],
                points=points,
                bbox=(min(xs), min(ys), max(xs), max(ys)),
                area=polygon_area(points),
            )
        )
    return polygons


def collect_labels(
    objects: list[dict[str, Any]], lookup: dict[int, str], layer: str
) -> list[Label]:
    labels: list[Label] = []
    for item in objects:
        if item.get("entity") != "TEXT" or object_layer(item, lookup) != layer:
            continue
        point = item.get("ins_pt")
        if not point or item.get("text_value") in (None, ""):
            continue
        labels.append(
            Label(
                object_index=item["index"],
                text=str(item["text_value"]),
                normalized=normalized_survey(item["text_value"]),
                point=(float(point[0]), float(point[1])),
            )
        )
    return labels


def assign_labels_to_polygons(labels: list[Label], polygons: list[Polygon]) -> tuple[dict[int, dict[str, Any]], Counter]:
    """Choose the smallest containing polygon if CAD contains duplicate outlines."""

    assignments: dict[int, dict[str, Any]] = {}
    candidate_counts: Counter = Counter()
    for label in labels:
        candidates = [
            polygon
            for polygon in polygons
            if (
                polygon.bbox[0] <= label.point[0] <= polygon.bbox[2]
                and polygon.bbox[1] <= label.point[1] <= polygon.bbox[3]
                and contains_point(label.point, polygon.points)
            )
        ]
        candidate_counts[len(candidates)] += 1
        if not candidates:
            assignments[label.object_index] = {"label": label, "polygon": None, "candidate_count": 0}
            continue
        selected = min(candidates, key=lambda polygon: (polygon.area, polygon.object_index))
        assignments[label.object_index] = {
            "label": label,
            "polygon": selected,
            "candidate_count": len(candidates),
        }
    return assignments, candidate_counts


def load_workbook_records(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        if worksheet.title not in {"Bhavanipar", "Bitta", "Vandh Timbo"}:
            continue
        headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        try:
            survey_column = headers.index("New Survey No.")
        except ValueError as exc:
            raise RuntimeError(f"New Survey No. column missing in {worksheet.title}") from exc
        consent_column = next(
            (
                index
                for index, header in enumerate(headers)
                if normalized_survey(header) in {"consent", "concentrecevied"}
            ),
            None,
        )
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            raw_survey = row[survey_column]
            if raw_survey in (None, ""):
                continue
            normalized = normalized_survey(raw_survey)
            consent_value = row[consent_column] if consent_column is not None else None
            records.append(
                {
                    "village": worksheet.title,
                    "survey_number": str(raw_survey).strip(),
                    "survey_normalized": normalized,
                    "survey_base": survey_base(normalized),
                    "source_row": row_number,
                    "consent_status": "received"
                    if normalized_survey(consent_value) == "yes"
                    else "not_recorded",
                }
            )
    return records


def workbook_indexes(records: list[dict[str, Any]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    exact: dict[str, list[dict[str, Any]]] = defaultdict(list)
    base: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        exact[record["survey_normalized"]].append(record)
        base[record["survey_base"]].append(record)
    return exact, base


def serialize_binding(record: dict[str, Any], match_kind: str, is_primary_village: bool) -> dict[str, Any]:
    return {
        "village": record["village"],
        "survey_number": record["survey_number"],
        "source_row": record["source_row"],
        "source_consent_status": record["consent_status"],
        "match_kind": match_kind,
        "requires_village_validation": not is_primary_village,
    }


def make_feature_bindings(
    label: Label,
    exact_index: dict[str, list[dict[str, Any]]],
    base_index: dict[str, list[dict[str, Any]]],
    primary_village: str,
) -> tuple[list[dict[str, Any]], str, str]:
    exact = exact_index.get(label.normalized, [])
    match_kind = "exact" if exact else "base_candidate"
    candidates = exact or base_index.get(survey_base(label.normalized), [])
    primary_records = [record for record in candidates if record["village"] == primary_village]
    bindings = [
        serialize_binding(record, match_kind, record["village"] == primary_village)
        for record in candidates
    ]
    if primary_records and all(record["consent_status"] == "received" for record in primary_records):
        return bindings, "received", "bhavanipar_candidate"
    if primary_records:
        return bindings, "unclassified", "bhavanipar_candidate"
    if bindings:
        return bindings, "unclassified", "cross_village_candidate"
    return bindings, "unclassified", "no_workbook_candidate"


def svg_element_index(root: ET.Element) -> dict[str, ET.Element]:
    return {
        element.attrib["id"]: element
        for element in root.iter()
        if element.attrib.get("id", "").startswith("dwg-object-")
    }


def svg_coordinate_origin(source_root: ET.Element, reference_polygon: Polygon) -> tuple[float, float]:
    """Derive the CAD-world to SVG-local offset from one known CAD path."""

    source_path = svg_element_index(source_root).get(
        f"dwg-object-{reference_polygon.object_index}"
    )
    if source_path is None:
        raise RuntimeError("Could not find a primary CAD polygon in the SVG export.")
    match = re.search(
        r"\bM\s*([-+0-9.eE]+),\s*([-+0-9.eE]+)", source_path.attrib.get("d", "")
    )
    if match is None:
        raise RuntimeError("Could not read the first point from the SVG CAD path.")
    local_x, local_y = (float(value) for value in match.groups())
    world_x, world_y = reference_polygon.points[0]
    return world_x - local_x, world_y - local_y


def fitted_view_box(
    features: list[dict[str, Any]], origin: tuple[float, float], fallback: str
) -> str:
    """Fit the browser view to the selected survey polygons instead of empty CAD space."""

    bboxes = [feature["cad_global_bbox"] for feature in features]
    if not bboxes:
        return fallback
    min_x = min(bbox[0] for bbox in bboxes) - origin[0]
    min_y = min(bbox[1] for bbox in bboxes) - origin[1]
    max_x = max(bbox[2] for bbox in bboxes) - origin[0]
    max_y = max(bbox[3] for bbox in bboxes) - origin[1]
    padding = max(max_x - min_x, max_y - min_y) * 0.04
    return f"{min_x - padding:g} {min_y - padding:g} {max_x - min_x + 2 * padding:g} {max_y - min_y + 2 * padding:g}"


def clean_view_box(value: str) -> str:
    values = value.replace(",", " ").split()
    if len(values) != 4:
        raise RuntimeError(f"Unexpected SVG viewBox: {value!r}")
    x, y, width, height = (float(number) for number in values)
    # LibreDWG sometimes applies a world-coordinate origin to a drawing whose
    # SVG path coordinates are already local. Preserve a known-good local origin.
    if abs(x) > 100_000 or abs(y) > 100_000:
        x, y = 0.0, 0.0
    return f"{x:g} {y:g} {width:g} {height:g}"


def element_with_attributes(element: ET.Element, allowed: set[str]) -> ET.Element:
    return ET.Element(
        element.tag,
        {key: value for key, value in element.attrib.items() if key in allowed},
    )


def build_svg(
    source_root: ET.Element,
    features: list[dict[str, Any]],
    label_lookup: dict[int, Label],
    view_box: str,
) -> ET.Element:
    source_elements = svg_element_index(source_root)
    root = ET.Element(
        f"{{{SVG_NS}}}svg",
        {
            "version": "1.1",
            "viewBox": view_box,
            "role": "img",
            "aria-labelledby": "map-title map-description",
            "preserveAspectRatio": "xMidYMid meet",
        },
    )
    title = ET.SubElement(root, f"{{{SVG_NS}}}title", {"id": "map-title"})
    title.text = "Bhachunda Solar combined village survey map"
    description = ET.SubElement(root, f"{{{SVG_NS}}}desc", {"id": "map-description"})
    description.text = (
        "Interactive CAD-derived survey map. Green shows only a received consent "
        "candidate from the Bhavanipar source sheet; grey shapes are not confirmed."
    )
    style = ET.SubElement(root, f"{{{SVG_NS}}}style")
    style.text = """
      .map-background { fill: #f5faf7; }
      .parcel-feature { fill: #e7eeea; stroke: #71877d; stroke-width: 0.8; cursor: pointer; }
      .parcel-feature:hover, .parcel-feature:focus { fill: #bfd7c9; stroke: #12352b; stroke-width: 1.8; outline: none; }
      .parcel-feature.consent-received { fill: #45b96f; stroke: #17633d; }
      .parcel-label { fill: #26433a; font-family: ui-monospace, SFMono-Regular, Menlo, monospace; font-size: 23px; pointer-events: none; user-select: none; }
    """
    view_box_values = view_box.split()
    ET.SubElement(root, f"{{{SVG_NS}}}rect", {
        "class": "map-background",
        "x": view_box_values[0],
        "y": view_box_values[1],
        "width": view_box_values[2],
        "height": view_box_values[3],
    })
    paths_group = ET.SubElement(root, f"{{{SVG_NS}}}g", {"id": "primary-parcel-boundaries"})
    labels_group = ET.SubElement(root, f"{{{SVG_NS}}}g", {"id": "primary-survey-labels"})

    for feature in features:
        source_path = source_elements.get(feature["svg_element_id"])
        label = label_lookup[feature["label_object_id"]]
        source_label = source_elements.get(f"dwg-object-{label.object_index}")
        if source_path is None or source_label is None:
            raise RuntimeError(
                f"SVG object missing for {feature['svg_element_id']} / dwg-object-{label.object_index}"
            )
        path = element_with_attributes(source_path, {"id", "d"})
        path.attrib.update(
            {
                "class": f"parcel-feature consent-{feature['display_consent_status']}",
                "data-feature-key": feature["feature_key"],
                "data-cad-survey": feature["cad_survey_number"],
                "data-display-consent": feature["display_consent_status"],
                "data-workbook-binding": feature["workbook_binding_state"],
                "data-linked-record-count": str(feature["primary_village_record_count"]),
                "tabindex": "0",
                "role": "button",
                "aria-label": feature["aria_label"],
            }
        )
        paths_group.append(path)

        text = element_with_attributes(source_label, {"id", "x", "y", "font-size"})
        text.attrib["class"] = "parcel-label"
        text.text = label.text
        labels_group.append(text)
    return root


def group_workbook_summary(records: list[dict[str, Any]], features: list[dict[str, Any]], primary_village: str) -> dict[str, Any]:
    summary: dict[str, Any] = {}
    for village in ("Bhavanipar", "Bitta", "Vandh Timbo"):
        village_records = [record for record in records if record["village"] == village]
        matched_rows = {
            (binding["village"], binding["source_row"])
            for feature in features
            for binding in feature["workbook_candidates"]
            if binding["village"] == village
        }
        received_rows = sum(record["consent_status"] == "received" for record in village_records)
        summary[village] = {
            "source_record_count": len(village_records),
            "received_source_record_count": received_rows,
            "candidate_source_record_count": len(matched_rows),
            "display_status_applied": village == primary_village,
        }
    summary[primary_village]["received_feature_count"] = sum(
        feature["display_consent_status"] == "received" for feature in features
    )
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-json", type=Path, required=True)
    parser.add_argument("--source-svg", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-svg", type=Path, required=True)
    parser.add_argument("--output-features", type=Path, required=True)
    parser.add_argument("--primary-village", default="Bhavanipar")
    args = parser.parse_args()

    drawing = json.loads(args.source_json.read_text(encoding="utf-8"))
    objects = drawing["OBJECTS"]
    layers = layer_name_lookup(objects)
    primary_polygons = collect_polygons(objects, layers, PRIMARY_POLYGON_LAYER)
    primary_labels = collect_labels(objects, layers, PRIMARY_LABEL_LAYER)
    primary_assignments, primary_counts = assign_labels_to_polygons(primary_labels, primary_polygons)

    secondary_polygons = collect_polygons(objects, layers, SECONDARY_POLYGON_LAYER)
    secondary_labels = collect_labels(objects, layers, SECONDARY_LABEL_LAYER)
    _, secondary_counts = assign_labels_to_polygons(secondary_labels, secondary_polygons)

    if any(result["polygon"] is None for result in primary_assignments.values()):
        raise RuntimeError("At least one primary CAD survey label could not be matched to a polygon.")

    source_root = ET.parse(args.source_svg).getroot()
    source_origin = svg_coordinate_origin(source_root, primary_polygons[0])
    records = load_workbook_records(args.workbook)
    exact_index, base_index = workbook_indexes(records)
    label_lookup = {label.object_index: label for label in primary_labels}
    features: list[dict[str, Any]] = []
    for result in primary_assignments.values():
        label: Label = result["label"]
        polygon: Polygon = result["polygon"]
        bindings, display_status, binding_state = make_feature_bindings(
            label, exact_index, base_index, args.primary_village
        )
        primary_records = [
            binding for binding in bindings if binding["village"] == args.primary_village
        ]
        feature_key = f"combined-villages-primary-{polygon.object_index}"
        feature = {
            "feature_key": feature_key,
            "svg_element_id": f"dwg-object-{polygon.object_index}",
            "source_cad_layer": PRIMARY_POLYGON_LAYER,
            "cad_survey_number": label.text,
            "cad_survey_normalized": label.normalized,
            "label_object_id": label.object_index,
            "match_method": "label_point_containment",
            "geometry_candidate_count": result["candidate_count"],
            "cad_global_bbox": [round(value, 6) for value in polygon.bbox],
            "workbook_binding_state": binding_state,
            "primary_village_record_count": len(primary_records),
            "display_consent_status": display_status,
            "workbook_candidates": bindings,
            "aria_label": (
                f"Survey {label.text}. "
                + (
                    "Consent received candidate from the Bhavanipar source sheet."
                    if display_status == "received"
                    else "Consent not yet confirmed on this CAD shape."
                )
            ),
        }
        features.append(feature)

    features.sort(key=lambda feature: int(feature["svg_element_id"].rsplit("-", 1)[-1]))
    view_box = fitted_view_box(
        features,
        source_origin,
        clean_view_box(source_root.attrib.get("viewBox", "")),
    )
    output_svg = build_svg(source_root, features, label_lookup, view_box)
    ET.indent(output_svg, space="  ")
    args.output_svg.parent.mkdir(parents=True, exist_ok=True)
    args.output_svg.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(output_svg, encoding="unicode"),
        encoding="utf-8",
    )

    manifest = {
        "schema_version": 1,
        "status": "staging_review_required",
        "source": {
            "cad_json": args.source_json.name,
            "cad_svg": args.source_svg.name,
            "workbook": args.workbook.name,
            "primary_village_for_preview": args.primary_village,
        },
        "cad_geometry": {
            "primary": {
                "polygon_layer": PRIMARY_POLYGON_LAYER,
                "label_layer": PRIMARY_LABEL_LAYER,
                "polygon_count": len(primary_polygons),
                "label_count": len(primary_labels),
                "label_to_polygon_candidate_counts": dict(sorted(primary_counts.items())),
            },
            "secondary": {
                "polygon_layer": SECONDARY_POLYGON_LAYER,
                "label_layer": SECONDARY_LABEL_LAYER,
                "polygon_count": len(secondary_polygons),
                "label_count": len(secondary_labels),
                "label_to_polygon_candidate_counts": dict(sorted(secondary_counts.items())),
            },
        },
        "workbook_coverage": group_workbook_summary(records, features, args.primary_village),
        "features": features,
    }
    args.output_features.parent.mkdir(parents=True, exist_ok=True)
    args.output_features.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(
        "Generated "
        f"{len(features)} primary map features; "
        f"geometry label candidates {dict(sorted(primary_counts.items()))}; "
        f"green preview features {sum(f['display_consent_status'] == 'received' for f in features)}."
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (RuntimeError, OSError, ValueError, KeyError) as error:
        print(f"Map build failed: {error}", file=sys.stderr)
        raise SystemExit(1)
