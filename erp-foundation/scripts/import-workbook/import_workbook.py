#!/usr/bin/env python3
"""Idempotent staging importer for the Bhachunda Solar ERP workbooks.

The default mode is a dry run. --apply requires SUPABASE_URL and
SUPABASE_SERVICE_ROLE_KEY in the environment. It never writes to the source
workbooks and deliberately never treats spreadsheet colour or a text link as a
Google Drive upload.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


DIGIT_TRANSLATION = str.maketrans("૦૧૨૩૪૫૬૭૮૯", "0123456789")
LAND_SHEETS = {
    "Bhavanipar": ("bhavanipar", "Bhavanipar"),
    "Bitta": ("bitta", "Bitta"),
    "Vandh Timbo": ("vandh-timbo", "Vandh Timbo"),
}
VILLAGE_ALIASES = {
    "bhavanipar": "bhavanipar",
    "bhavani par": "bhavanipar",
    "bitta": "bitta",
    "vandh timbo": "vandh-timbo",
    "vandh timba": "vandh-timbo",
}


class ImportErrorSafe(RuntimeError):
    """A failure whose message is safe to print without source data."""


@dataclass
class Summary:
    land_rows: int = 0
    patel_rows: int = 0
    matched_patel_rows: int = 0
    parcels_upserted: int = 0
    owner_rows_upserted: int = 0
    consent_rows_upserted: int = 0
    cases_upserted: int = 0
    legal_rows_upserted: int = 0
    private_rows_upserted: int = 0
    skipped: Counter[str] = field(default_factory=Counter)
    village_totals: Counter[str] = field(default_factory=Counter)
    received_totals: Counter[str] = field(default_factory=Counter)

    def show(self) -> None:
        print("\nImport reconciliation")
        print(f"  Land worksheet rows accepted: {self.land_rows}")
        print(f"  Patel Infra rows read: {self.patel_rows}")
        print(f"  Patel Infra rows matched: {self.matched_patel_rows}")
        if self.parcels_upserted:
            print(f"  Parcels upserted: {self.parcels_upserted}")
            print(f"  Owner rows upserted: {self.owner_rows_upserted}")
            print(f"  Consent rows upserted: {self.consent_rows_upserted}")
            print(f"  Acquisition rows upserted: {self.cases_upserted}")
            print(f"  Legal rows upserted: {self.legal_rows_upserted}")
            if self.private_rows_upserted:
                print(f"  Private owner rows upserted: {self.private_rows_upserted}")
        print("  Village controls:")
        for code in ("bhavanipar", "bitta", "vandh-timbo"):
            print(f"    {code}: {self.village_totals[code]} surveys, {self.received_totals[code]} consent received")
        if self.skipped:
            print("  Exceptions (counts only):")
            for reason, count in sorted(self.skipped.items()):
                print(f"    {reason}: {count}")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalise_header(value: Any) -> str:
    return re.sub(r"\s+", " ", text(value).replace("\n", " ").lower()).strip()


def normalise_survey(value: Any) -> str:
    return re.sub(r"\s+", "", text(value).translate(DIGIT_TRANSLATION).lower())


def normalise_village(value: Any) -> str | None:
    candidate = re.sub(r"\s+", " ", text(value).lower()).strip()
    return VILLAGE_ALIASES.get(candidate)


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    candidate = text(value).translate(DIGIT_TRANSLATION).replace(",", "")
    if candidate.lower() in {"", "na", "n/a", "not available", "-"}:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", candidate)
    if not match:
        return None
    try:
        return float(Decimal(match.group(0)))
    except InvalidOperation:
        return None


def whole_number(value: Any) -> int | None:
    parsed = number(value)
    return int(parsed) if parsed is not None else None


def iso_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%b-%Y", "%B-%Y"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            continue
    return None


def json_value(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return iso_date(value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return text(value)


def header_index(row: Iterable[Any]) -> dict[str, int]:
    return {normalise_header(cell): index for index, cell in enumerate(row) if text(cell)}


def find_cell(row: tuple[Any, ...], headers: dict[str, int], *required: str, exclude: tuple[str, ...] = ()) -> Any:
    for header, index in headers.items():
        if all(term in header for term in required) and not any(term in header for term in exclude):
            return row[index] if index < len(row) else None
    return None


def status_from_consent(value: Any) -> str:
    raw = text(value).lower()
    if raw in {"yes", "y", "received", "done", "1"}:
        return "received"
    if raw in {"pending", "p", "in progress"}:
        return "pending"
    if raw in {"blocked", "hold"}:
        return "blocked"
    if raw in {"rejected", "no"}:
        return "rejected"
    return "not_ready"


def owner_parts(value: Any) -> list[str]:
    original = text(value)
    if not original:
        return []
    # Land workbooks number owners on separate lines. Preserve an unparsed
    # multi-line block as one owner rather than accidentally splitting a name.
    sections = re.split(r"\n(?=\s*\d+[.)])", original)
    parsed = [re.sub(r"^\s*\d+[.)]\s*", "", section).strip() for section in sections]
    return [part for part in parsed if part]


def compact_payload(data: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in data.items() if value is not None and value != ""}


def parse_land_workbook(path: Path, summary: Summary) -> list[dict[str, Any]]:
    if not path.exists():
        raise ImportErrorSafe("Land workbook path does not exist.")
    records: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name, (village_code, village_name) in LAND_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                summary.skipped[f"missing_sheet_{sheet_name.lower().replace(' ', '_')}"] += 1
                continue
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headings = next(iterator, None)
            if not headings:
                summary.skipped[f"empty_sheet_{sheet_name.lower().replace(' ', '_')}"] += 1
                continue
            headers = header_index(headings)
            seen: set[str] = set()
            for source_row, row in enumerate(iterator, start=2):
                survey = find_cell(row, headers, "new survey")
                survey_normalised = normalise_survey(survey)
                if not survey_normalised:
                    has_land_record = any((
                        text(find_cell(row, headers, "farmer name")),
                        text(find_cell(row, headers, "account number")),
                        text(find_cell(row, headers, "old survey")),
                    ))
                    if not has_land_record:
                        # Excel's used range contains blank trailing rows and
                        # acreage-only total rows; neither is a land record.
                        continue
                    summary.skipped["blank_survey_number"] += 1
                    continue
                if survey_normalised in seen:
                    summary.skipped["duplicate_normalised_survey_in_source"] += 1
                    continue
                seen.add(survey_normalised)
                consent_raw = find_cell(row, headers, "consent")
                # The source typo is "Concent Recevied" in Bhavanipar. It is
                # intentionally covered by the second lookup, without reading fill colours.
                if consent_raw is None:
                    consent_raw = find_cell(row, headers, "concent")
                record = {
                    "village_code": village_code,
                    "village_name": village_name,
                    "village_name_gu": text(find_cell(row, headers, "village")) or None,
                    "district": text(find_cell(row, headers, "district")) or None,
                    "taluka": text(find_cell(row, headers, "taluka")) or None,
                    "survey_number": text(survey),
                    "survey_normalised": survey_normalised,
                    "old_survey_number": text(find_cell(row, headers, "old survey")) or None,
                    "hectare_are_sqmt": text(find_cell(row, headers, "h.are.sqmt")) or None,
                    "acreage": number(find_cell(row, headers, "acer guntha")),
                    "account_number": text(find_cell(row, headers, "account number")) or None,
                    "owner_text": text(find_cell(row, headers, "farmer name")) or None,
                    "tenure": text(find_cell(row, headers, "tenure")) or None,
                    "land_use": text(find_cell(row, headers, "land use")) or None,
                    "rights_and_encumbrances": text(find_cell(row, headers, "હક્ક")) or None,
                    "nondh_numbers": text(find_cell(row, headers, "nondh number")) or None,
                    "bunch_number": text(find_cell(row, headers, "bunch number")) or None,
                    "consent_raw": text(consent_raw) or None,
                    "source_row_number": source_row,
                    "source_workbook": path.name,
                }
                records.append(record)
                summary.land_rows += 1
                summary.village_totals[village_code] += 1
                if status_from_consent(record["consent_raw"]) == "received":
                    summary.received_totals[village_code] += 1
    finally:
        workbook.close()
    return records


def parse_patel_workbook(path: Path, summary: Summary) -> list[dict[str, Any]]:
    if not path.exists():
        raise ImportErrorSafe("Patel Infra workbook path does not exist.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headings = next(iterator, None)
        if not headings:
            raise ImportErrorSafe("Patel Infra workbook does not have a header row.")
        headers = header_index(headings)
        # Row 2 is the template's source/system label and intentionally skipped.
        for source_row, row in enumerate(iterator, start=2):
            if source_row == 2:
                continue
            survey = find_cell(row, headers, "survey no", exclude=("unique",))
            village = normalise_village(find_cell(row, headers, "village"))
            if not text(survey) and not village:
                continue
            if not village or not normalise_survey(survey):
                summary.skipped["patel_row_missing_village_or_survey"] += 1
                continue
            values = {header: row[index] if index < len(row) else None for header, index in headers.items()}
            records.append({
                "village_code": village,
                "survey_normalised": normalise_survey(survey),
                "source_row_number": source_row,
                "values": values,
            })
            summary.patel_rows += 1
    finally:
        workbook.close()
    return records


class RestClient:
    def __init__(self, project_url: str, service_role_key: str):
        self.base_url = project_url.rstrip("/")
        self.service_role_key = service_role_key

    def request(self, method: str, table: str, *, payload: Any = None, query: dict[str, str] | None = None, prefer: str | None = None) -> Any:
        url = f"{self.base_url}/rest/v1/{quote(table, safe='_')}"
        if query:
            url += "?" + urlencode(query, safe="(),.*")
        headers = {
            "apikey": self.service_role_key,
            "Authorization": f"Bearer {self.service_role_key}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        }
        if prefer:
            headers["Prefer"] = prefer
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8") if payload is not None else None
        request = Request(url, data=body, headers=headers, method=method)
        try:
            with urlopen(request, timeout=60) as result:
                content = result.read().decode("utf-8")
                return json.loads(content) if content else None
        except HTTPError as error:
            raise ImportErrorSafe(f"Supabase REST request failed for {table} ({error.code}).") from error
        except URLError as error:
            raise ImportErrorSafe("Could not reach Supabase. Check SUPABASE_URL and network access.") from error

    def upsert(self, table: str, payload: dict[str, Any], conflict: str) -> dict[str, Any]:
        result = self.request(
            "POST", table, payload=payload, query={"on_conflict": conflict},
            prefer="resolution=merge-duplicates,return=representation",
        )
        if not isinstance(result, list) or not result:
            raise ImportErrorSafe(f"Supabase did not return an upserted {table} row.")
        return result[0]

    def insert(self, table: str, payload: dict[str, Any]) -> dict[str, Any]:
        result = self.request("POST", table, payload=payload, prefer="return=representation")
        if not isinstance(result, list) or not result:
            raise ImportErrorSafe(f"Supabase did not return an inserted {table} row.")
        return result[0]

    def patch(self, table: str, payload: dict[str, Any], query: dict[str, str]) -> None:
        self.request("PATCH", table, payload=payload, query=query, prefer="return=minimal")

    def select(self, table: str, query: dict[str, str]) -> list[dict[str, Any]]:
        result = self.request("GET", table, query=query)
        return result if isinstance(result, list) else []


def checksum(paths: list[Path]) -> str:
    digest = hashlib.sha256()
    for path in paths:
        with path.open("rb") as source:
            for block in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(block)
    return digest.hexdigest()


def patel_value(values: dict[str, Any], *required: str, exclude: tuple[str, ...] = ()) -> Any:
    for header, value in values.items():
        if all(term in header for term in required) and not any(term in header for term in exclude):
            return value
    return None


def patch_patel_record(api: RestClient, parcel_id: str, values: dict[str, Any], include_private: bool, summary: Summary) -> None:
    latitude = number(patel_value(values, "latitude"))
    longitude = number(patel_value(values, "longitude"))
    if latitude is not None or longitude is not None:
        api.patch("parcels", compact_payload({"latitude": latitude, "longitude": longitude}), {"id": f"eq.{parcel_id}"})

    source_fields = compact_payload({
        "state": json_value(patel_value(values, "state")),
        "district": json_value(patel_value(values, "district")),
        "tehsil": json_value(patel_value(values, "tehsil")),
        "mamlatdar_office": json_value(patel_value(values, "mamlatdar")),
        "sub_registrar_office": json_value(patel_value(values, "sub registrar")),
        "sdm_office": json_value(patel_value(values, "sdm office")),
        "dc_office": json_value(patel_value(values, "dc office")),
        "patwari_details": json_value(patel_value(values, "patwari")),
        "tehsildar_details": json_value(patel_value(values, "tehsildar")),
        "police_station": json_value(patel_value(values, "police station")),
        "sp_office": json_value(patel_value(values, "sp office")),
        "aggregator_name": json_value(patel_value(values, "aggregator name")),
        "aggregator_rep": json_value(patel_value(values, "aggregator rep")),
        "aging": json_value(patel_value(values, "aging")),
    })
    acquisition = compact_payload({
        "parcel_id": parcel_id,
        "project_name": text(patel_value(values, "project name")) or None,
        "spv_name": text(patel_value(values, "spv name")) or None,
        "mw": number(patel_value(values, "mw")),
        "category": text(patel_value(values, "category", exclude=("atl",))) or None,
        "atl_category": text(patel_value(values, "category", "atl")) or None,
        "acquisition_purpose": text(patel_value(values, "acquisition purpose")) or None,
        "project_duration_months": whole_number(patel_value(values, "project duration")),
        "block_name": text(patel_value(values, "block")) or None,
        "target_date": iso_date(patel_value(values, "target date")),
        "execution_date": iso_date(patel_value(values, "date of execution")),
        "reason_not_acquired": text(patel_value(values, "reasons for not acquired")) or None,
        "total_acres_in_rtc": number(patel_value(values, "total acres in rtc")),
        "total_acres_to_acquire": number(patel_value(values, "total to acquired")),
        "total_acres_acquired": number(patel_value(values, "total acres aquired")),
        "total_sq_meters_acquired": number(patel_value(values, "total sq. mtrs")),
        "source_fields": source_fields,
    })
    api.upsert("acquisition_cases", acquisition, "parcel_id")
    summary.cases_upserted += 1

    legal = compact_payload({
        "parcel_id": parcel_id,
        "public_notice_status": text(patel_value(values, "public notice")) or None,
        "sro_search_status": text(patel_value(values, "sro search")) or None,
        "documents_required": whole_number(patel_value(values, "required no.of documents")),
        "documents_submitted": whole_number(patel_value(values, "submited no.of documents")),
        "law_firm_verification_status": text(patel_value(values, "verification by law firm")) or None,
        "pending_documents": text(patel_value(values, "pending documents")) or None,
        "preliminary_tsr_status": text(patel_value(values, "preliminary tsr")) or None,
        "conditional_clearance_status": text(patel_value(values, "conditional cleareance")) or None,
        "nfa_number": text(patel_value(values, "nfa no")) or None,
        "nfa_submitted_on": iso_date(patel_value(values, "nfa submission")),
        "nfa_approved_on": iso_date(patel_value(values, "nfa approval")),
        "legal_remarks": text(patel_value(values, "remarks")) or None,
    })
    api.upsert("legal_reviews", legal, "parcel_id")
    summary.legal_rows_upserted += 1

    if not include_private:
        return
    private = compact_payload({
        "pan_number": text(patel_value(values, "pan card")) or None,
        "aadhaar_number": text(patel_value(values, "aadhar")) or None,
        "bank_account_number": text(patel_value(values, "bank account")) or None,
        "bank_name": text(patel_value(values, "bank name")) or None,
        "ifsc_code": text(patel_value(values, "ifsc")) or None,
        "vendor_code": text(patel_value(values, "vendor code")) or None,
        "bank_owner_name": text(patel_value(values, "owner name")) or text(patel_value(values, "owners name", "bank records")) or None,
    })
    if not private:
        return
    owner_rows = api.select("parcel_owners", {"select": "id", "parcel_id": f"eq.{parcel_id}", "order": "sequence_no.asc", "limit": "1"})
    if owner_rows:
        owner_id = owner_rows[0]["id"]
    else:
        fallback_name = private.get("bank_owner_name") or "Owner awaiting reconciliation"
        owner = api.upsert("parcel_owners", {"parcel_id": parcel_id, "display_name": fallback_name, "sequence_no": 1, "is_primary": True}, "parcel_id,sequence_no")
        owner_id = owner["id"]
        summary.owner_rows_upserted += 1
    api.upsert("owner_private_details", {"owner_id": owner_id, **private}, "owner_id")
    summary.private_rows_upserted += 1


def run_apply(land_records: list[dict[str, Any]], patel_records: list[dict[str, Any]], land_path: Path, patel_path: Path, include_private: bool, summary: Summary) -> None:
    project_url = os.environ.get("SUPABASE_URL", "").strip()
    service_role_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    if not project_url or not service_role_key:
        raise ImportErrorSafe("--apply requires SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY in the environment.")
    api = RestClient(project_url, service_role_key)
    run = api.insert("import_runs", {
        "source_name": f"{land_path.name}; {patel_path.name}",
        "source_checksum": checksum([land_path, patel_path]),
        "notes": "Staging workbook import started by import_workbook.py",
    })
    village_ids: dict[str, str] = {}
    parcel_index: dict[tuple[str, str], str] = {}
    try:
        for record in land_records:
            code = record["village_code"]
            if code not in village_ids:
                village = api.upsert("villages", compact_payload({
                    "code": code,
                    "name_en": record["village_name"],
                    "name_gu": record["village_name_gu"],
                    "district": record["district"],
                    "taluka": record["taluka"],
                }), "code")
                village_ids[code] = village["id"]
            parcel = api.upsert("parcels", compact_payload({
                "village_id": village_ids[code],
                "survey_number": record["survey_number"],
                "old_survey_number": record["old_survey_number"],
                "hectare_are_sqmt": record["hectare_are_sqmt"],
                "acreage": record["acreage"],
                "account_number": record["account_number"],
                "tenure": record["tenure"],
                "land_use": record["land_use"],
                "rights_and_encumbrances": record["rights_and_encumbrances"],
                "nondh_numbers": record["nondh_numbers"],
                "bunch_number": record["bunch_number"],
                "source_workbook": record["source_workbook"],
                "source_row_number": record["source_row_number"],
            }), "village_id,survey_number_normalized")
            parcel_id = parcel["id"]
            parcel_index[(code, record["survey_normalised"])] = parcel_id
            summary.parcels_upserted += 1
            for sequence, name in enumerate(owner_parts(record["owner_text"]), start=1):
                api.upsert("parcel_owners", {
                    "parcel_id": parcel_id,
                    "display_name": name,
                    "source_owner_text": record["owner_text"],
                    "sequence_no": sequence,
                    "is_primary": sequence == 1,
                }, "parcel_id,sequence_no")
                summary.owner_rows_upserted += 1
            consent_status = status_from_consent(record["consent_raw"])
            api.upsert("consent_records", {
                "parcel_id": parcel_id,
                "status": consent_status,
                "source_value": record["consent_raw"],
            }, "parcel_id")
            summary.consent_rows_upserted += 1

        for record in patel_records:
            parcel_id = parcel_index.get((record["village_code"], record["survey_normalised"]))
            if not parcel_id:
                summary.skipped["patel_survey_not_matched_to_land_workbook"] += 1
                continue
            patch_patel_record(api, parcel_id, record["values"], include_private, summary)
            summary.matched_patel_rows += 1

        api.patch("import_runs", {
            "completed_at": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
            "inserted_count": summary.parcels_upserted,
            "updated_count": summary.cases_upserted + summary.legal_rows_upserted,
            "skipped_count": sum(summary.skipped.values()),
            "error_count": 0,
            "notes": "Completed; review reconciliation counts before production use.",
        }, {"id": f"eq.{run['id']}"})
    except Exception:
        # The upserts are safe to rerun. Record that the owner must review this
        # run instead of reporting a misleading completed state.
        try:
            api.patch("import_runs", {
                "completed_at": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
                "skipped_count": sum(summary.skipped.values()),
                "error_count": 1,
                "notes": "Stopped with an error; rerun after correcting the reported issue.",
            }, {"id": f"eq.{run['id']}"})
        except Exception:
            pass
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description="Reconcile Bhachunda Solar workbooks into a staging Supabase project.")
    parser.add_argument("--land-workbook", required=True, type=Path, help="BHAVANIPAR AND BITTA FINAL SHEET workbook")
    parser.add_argument("--patel-workbook", required=True, type=Path, help="PATEL INFRA official land format workbook")
    parser.add_argument("--apply", action="store_true", help="Write approved data to Supabase. Dry run is the default.")
    parser.add_argument("--include-private", action="store_true", help="Also import PAN, Aadhaar and bank fields into the restricted private table. Requires --apply.")
    args = parser.parse_args()
    if args.include_private and not args.apply:
        parser.error("--include-private can be used only with --apply.")

    summary = Summary()
    land_records = parse_land_workbook(args.land_workbook, summary)
    patel_records = parse_patel_workbook(args.patel_workbook, summary)
    if not args.apply:
        # Count only in dry run; real upsert counters stay at zero to make the
        # output unmistakably non-mutating.
        land_keys = {(record["village_code"], record["survey_normalised"]) for record in land_records}
        summary.matched_patel_rows = sum((record["village_code"], record["survey_normalised"]) in land_keys for record in patel_records)
        summary.skipped["patel_survey_not_matched_to_land_workbook"] += len(patel_records) - summary.matched_patel_rows
        print("DRY RUN ONLY — no data was sent to Supabase.")
    else:
        run_apply(land_records, patel_records, args.land_workbook, args.patel_workbook, args.include_private, summary)
        print("APPLY COMPLETE — review the reconciliation below before using the production workflow.")
    summary.show()
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ImportErrorSafe as error:
        print(f"Import stopped: {error}", file=sys.stderr)
        raise SystemExit(2)
